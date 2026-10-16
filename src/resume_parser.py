"""Text extraction utilities for PDF, DOCX, and TXT files."""
from io import BytesIO
from pathlib import Path


def extract_text(filename: str, content: bytes) -> str:
    """Extract text from bytes based on file extension.

    Scanned/image-only PDFs may return little or no text because OCR is not included yet.
    """
    suffix = Path(filename).suffix.lower()

    if suffix == ".txt":
        return content.decode("utf-8", errors="replace")

    if suffix == ".pdf":
        import pymupdf as fitz
        pages = []
        with fitz.open(stream=content, filetype="pdf") as document:
            for page in document:
                pages.append(page.get_text())
        return "\n".join(pages).strip()

    if suffix == ".docx":
        from docx import Document
        document = Document(BytesIO(content))
        paragraphs = [p.text for p in document.paragraphs if p.text.strip()]
        table_text = []
        for table in document.tables:
            for row in table.rows:
                table_text.append(" | ".join(cell.text.strip() for cell in row.cells))
        return "\n".join(paragraphs + table_text).strip()

    raise ValueError("Unsupported file type. Please use PDF, DOCX, or TXT.")
